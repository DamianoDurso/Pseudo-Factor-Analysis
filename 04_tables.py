"""Step 5: results and manuscript tables, calculated from the step 4 outputs (nothing typed in).

Results.xlsx               one tab per study: factor alignment (from DAAL.csv); the loadings of every solution,
                           all three encodings side by side; congruence and correlations with the empirical benchmark;
                           the robustness check.
Manuscript Tables.xlsx     one tab per paper table, laid out as printed, and a tab of the numbers quoted in the text.

Summary cells (proportions, means, M, Overall M, totals) are Excel formulas over the cells they summarise.

As in the original analysis: loadings are the 3-decimal values; congruence (Tucker's phi, psych::fa.congruence) is
computed per aligned factor and rounded to 2 decimals; NA where the factor is not aligned; means are taken over the
2-decimal values. Each factor is oriented to agree with its empirical counterpart (absolute congruence). Printed values are rounded half away from zero, as Excel displays them.
"""
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parent

SOLUTIONS = ['avg_trans', 'distilroberta', 'miniLM', 'mpnet', 't5', 'psych', 'use_dan']
NAMES = {'avg_trans': 'Transformers', 'distilroberta': 'D_RoBERTa', 'miniLM': 'MiniLM', 'mpnet': 'MPNet',
         't5': 'T5', 'psych': 'Wulff & Mata', 'use_dan': 'USE-DAN'}
TABLE_ROWS = ['distilroberta', 'miniLM', 'mpnet', 't5', 'psych', 'use_dan', 'avg_trans']  # row order of the paper
ENCODINGS = {'item': 'Atomic', 'item_rev': 'Atomic reversed', 'sent': 'Macro'}

STUDIES = {
    'NEO': dict(factors=['O', 'C', 'E', 'A', 'N'], bench='NEO/benchmark_Johnson2014_Table4_left_panel.csv',
                source='Johnson, 2014, Table 4', perturb='NEOAC',
                abbrev=['Ima', 'Art', 'Emo', 'Adv', 'Int', 'Lib', 'Sel', 'Ord', 'Dut', 'Ach', 'Sel', 'Cau',
                        'Fri', 'Gre', 'Ass', 'Act', 'Exc', 'Che', 'Tru', 'Mor', 'Alt', 'Co', 'Mod', 'Sym',
                        'Anx', 'Ang', 'Dep', 'Sel', 'Imm', 'Vul']),
    'HEXACO': dict(factors=['H', 'E', 'X', 'A', 'C', 'O'], bench='HEXACO/benchmark_LeeAshton2004_Table4.csv',
                   source='Lee & Ashton, 2004, Table 4', perturb='HEXACO',
                   abbrev=None),  # HEXACO tables print the facet codes
}

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=12)
LINE = Side(style='thin')
key = lambda s: re.sub(r'[^a-z]', '', s.lower())


def r2(x):
    """Round to 2 decimals as Excel displays a number: 15 significant digits, then half away from zero."""
    return float(Decimal(f'{x:.15g}').quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def load(study):
    s = STUDIES[study]
    daal = pd.read_csv(ROOT / study / 'results' / 'DAAL.csv')
    loadings = {e: pd.read_csv(ROOT / study / 'results' / f'Results_target_{e}_labeled.csv', index_col=0)
                for e in ENCODINGS}
    bench = pd.read_csv(ROOT / s['bench'])
    full = dict(zip(bench['facet'].map(key) if study == 'NEO' else bench['code'], bench['facet']))
    facets = list(loadings['item'].index)
    s['full_names'] = [full[key(f)] if study == 'NEO' else full[f] for f in facets]
    s['codes'] = facets
    return daal, loadings


def solution_block(L, sol, factors):
    """Loadings of one solution, columns in theoretical order, headed by factor letters (merged: 'EA'; unassigned: 'U')."""
    cols = [c for c in L.columns if c.endswith('_' + sol) and not (sol == 'psych' and c.endswith('avg_trans'))]
    won = {c: re.findall(r'([A-Z])_' + sol, c) for c in cols if not c.startswith('ML')}
    def position(c):
        return factors.index(won[c][0]) if c in won else len(factors)
    cols = sorted(cols, key=position)
    return L[cols].set_axis([''.join(won[c]) if c in won else 'U' for c in cols], axis=1)


def alignment(daal, factors):
    """0/1 per factor (aligned under DAAL), rows = solutions in paper order, columns = encoding x factor."""
    a = daal.assign(aligned=daal['aligned'].astype(int)).pivot_table(
        index='solution', columns=['encoding', 'factor'], values='aligned')
    return a.reindex(index=TABLE_ROWS, columns=pd.MultiIndex.from_product([list(ENCODINGS), factors]))


def benchmark(study):
    """Empirical loadings, rows in the order of the results files, columns in theoretical factor order."""
    s = STUDIES[study]
    b = pd.read_csv(ROOT / s['bench'])
    b.index = b['facet'].map(key) if study == 'NEO' else b['code']
    return b.loc[[key(f) if study == 'NEO' else f for f in s['codes']], s['factors']].set_axis(s['codes'])


def congruence(study, loadings, bench):
    """Tucker's congruence of each aligned factor with its empirical factor (2 decimals, as psych::fa.congruence);
    rows = solutions in paper order plus 'Overall M', columns = encoding x factor plus 'M'."""
    f = STUDIES[study]['factors']
    rows = {}
    for sol in TABLE_ROWS:
        vals = []
        for e in ENCODINGS:
            b = solution_block(loadings[e], sol, f)
            for k in f:
                if k in b.columns:
                    x, y = bench[k].values, b[k].values.astype(float)
                    # factor signs are arbitrary: the factor is oriented to agree with its empirical counterpart
                    vals.append(round(abs(float(x @ y / np.sqrt((x @ x) * (y @ y)))), 2))
                else:
                    vals.append(np.nan)
        rows[NAMES[sol]] = vals
    c = pd.DataFrame(rows, index=pd.MultiIndex.from_product([list(ENCODINGS), f])).T
    c.loc['Overall M'] = c.mean()  # NaN-skipping column means over the 7 solutions
    c['M'] = c.mean(axis=1)        # row means over the non-NA cells
    return c


def correlations(study, loadings, bench):
    """Pearson correlations, over facets, of each empirical factor with each column of the averaged-transformer solution."""
    f = STUDIES[study]['factors']
    blocks = []
    for e, prefix in zip(ENCODINGS, 'ARO'):
        b = solution_block(loadings[e], 'avg_trans', f).astype(float)
        blocks.append(b.set_axis([prefix + h for h in b.columns], axis=1))
    emb = pd.concat(blocks, axis=1)
    return pd.DataFrame({c: [np.corrcoef(bench[k], emb[c])[0, 1] for k in f] for c in emb.columns},
                        index=['E' + k for k in f])


def printed_labels(columns, factors):
    """Correlation-table headings as printed in the paper: a merged column takes its first factor's letter and an
    unassigned column the letter of the factor left without a column (e.g. NEO atomic: AON -> AO, AU -> AN)."""
    out = []
    for prefix in 'ARO':
        group = [c[1:] for c in columns if c[0] == prefix]
        missing = iter([k for k in factors if k not in [g[0] for g in group if g != 'U']])
        out += [prefix + (g[0] if g != 'U' else next(missing)) for g in group]
    return out


def step_label(step, order):
    """'perturb_NE' -> 'N + E replaced', as printed in Tables 7 and 8."""
    if step == 'baseline':
        return 'Baseline'
    k = len(step) - len('perturb_')
    return 'All replaced' if k == len(order) else ' + '.join(order[:k]) + ' replaced'


def congruence_row(name, values, row, first, col0):
    """Cells of one congruence-table row starting in column col0: the solution's 2-decimal congruences (NA where
    not aligned) and the formula for M; for 'Overall M', formulas for the column means and the mean of the row."""
    n = len(values) - 1
    a, b = L(col0), L(col0 + n - 1)
    if name == 'Overall M':
        cells = [f'=IFERROR(AVERAGE({L(col0 + j)}{first}:{L(col0 + j)}{row - 1}),"NA")' for j in range(n)]
    else:
        cells = ['NA' if np.isnan(x) else r2(x) for x in values[:n]]
    return cells + [f'=IFERROR(AVERAGE({a}{row}:{b}{row}),"NA")']


def write_cell(ws, r, c, v, fmt=None, font=None):
    cell = ws.cell(r, c, v)
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font
    return cell


# ---------------------------------------------------------------- Results.xlsx
wb = Workbook()
wb.remove(wb.active)
data = {}
for study, s in STUDIES.items():
    daal, loadings = load(study)
    data[study] = (daal, loadings)
    f = s['factors']
    ws = wb.create_sheet(study)
    row = 1
    write_cell(ws, row, 1, f'{study}: factor alignment (DAAL; 1 = factor uniquely identified)', font=TITLE)
    row += 2
    al = alignment(daal, f)
    col = 2
    for e, label in ENCODINGS.items():
        write_cell(ws, row, col, label, font=BOLD)
        col += len(f)
    write_cell(ws, row, col, 'Proportion aligned', font=BOLD)
    row += 1
    write_cell(ws, row, 1, 'Solution', font=BOLD)
    heads = [x for e in ENCODINGS for x in f] + list(ENCODINGS.values()) + ['All']
    for j, h in enumerate(heads):
        write_cell(ws, row, 2 + j, h, font=BOLD)
    n, first = len(f), row + 1
    for sol in TABLE_ROWS:
        row += 1
        write_cell(ws, row, 1, NAMES[sol])
        for j, v in enumerate(al.loc[sol].astype(int)):
            write_cell(ws, row, 2 + j, int(v), fmt='0')
        for i in range(3):  # proportion aligned per encoding, then over all encodings
            write_cell(ws, row, 2 + 3 * n + i, f'=AVERAGE({L(2 + i * n)}{row}:{L(1 + (i + 1) * n)}{row})', fmt='0.000')
        write_cell(ws, row, 2 + 3 * n + 3, f'=AVERAGE(B{row}:{L(1 + 3 * n)}{row})', fmt='0.000')
    row += 1
    write_cell(ws, row, 1, 'Mean', font=BOLD)
    for j in range(3 * n + 4):
        write_cell(ws, row, 2 + j, f'=AVERAGE({L(2 + j)}{first}:{L(2 + j)}{row - 1})', fmt='0.000', font=BOLD)

    row += 3
    write_cell(ws, row, 1, f'{study}: target-rotated loadings (3 decimals). Column headings: aligned factor; '
                           'two letters = factors merged on one column; U = unassigned', font=TITLE)
    for sol in SOLUTIONS:
        row += 2
        nonconv = daal[(daal['solution'] == sol) & ~daal['rotation_converged']]['encoding'].unique()
        note = ('  (rotation did not converge: ' + ', '.join(ENCODINGS[e].lower() for e in nonconv) + ')'
                if len(nonconv) else '')
        write_cell(ws, row, 1, NAMES[sol] + note, font=BOLD)
        row += 1
        blocks = [solution_block(loadings[e], sol, f) for e in ENCODINGS]
        col = 2
        for b, label in zip(blocks, ENCODINGS.values()):
            write_cell(ws, row, col, label, font=BOLD)
            for j, h in enumerate(b.columns):
                write_cell(ws, row + 1, col + j, h, font=BOLD)
            for i, v in enumerate(b.values):
                for j, x in enumerate(v):
                    write_cell(ws, row + 2 + i, col + j, float(x), fmt='0.000')
            col += b.shape[1] + 1
        write_cell(ws, row + 1, 1, 'Facet', font=BOLD)
        for i, name in enumerate(s['full_names']):
            write_cell(ws, row + 2 + i, 1, name)
        row += 1 + len(s['full_names'])

    bench = benchmark(study)
    cong = congruence(study, loadings, bench)
    corr = correlations(study, loadings, bench)
    data[study] += (cong, corr)
    row += 2
    write_cell(ws, row, 1, f'{study}: congruence with the empirical loadings ({s["source"]}); '
                           'NA = factor not aligned', font=TITLE)
    row += 2
    col = 2
    for label in ENCODINGS.values():
        write_cell(ws, row, col, label, font=BOLD)
        col += len(f)
    row += 1
    write_cell(ws, row, 1, 'Solution', font=BOLD)
    for j, h in enumerate([k for _, k in cong.columns[:-1]] + ['M']):
        write_cell(ws, row, 2 + j, h, font=BOLD)
    for name, v in cong.iterrows():
        row += 1
        write_cell(ws, row, 1, name, font=BOLD if name == 'Overall M' else None)
        for j, cell in enumerate(congruence_row(name, v, row, row - len(TABLE_ROWS), 2)):
            write_cell(ws, row, 2 + j, cell, fmt='0.000')
    row += 3
    write_cell(ws, row, 1, f'{study}: correlations of the empirical loadings ({s["source"]}) with the averaged-'
                           'transformer loadings; A = atomic, R = atomic reversed, O = macro', font=TITLE)
    row += 2
    write_cell(ws, row, 1, 'Empirical', font=BOLD)
    for j, h in enumerate(corr.columns):
        write_cell(ws, row, 2 + j, h, font=BOLD)
    for name, v in corr.iterrows():
        row += 1
        write_cell(ws, row, 1, name)
        for j, x in enumerate(v):
            write_cell(ws, row, 2 + j, float(x), fmt='0.000')
    rob = pd.read_csv(ROOT / study / 'results' / 'robustness.csv')
    order = list(STUDIES[study]['perturb'])
    data[study] += (rob,)
    row += 3
    write_cell(ws, row, 1, f'{study}: robustness check, factors aligned (DAAL) as the embeddings of one factor '
                           'at a time are replaced by Gaussian noise; averaged-transformer matrix', font=TITLE)
    for floor in (0.30, 0.35):
        row += 2
        write_cell(ws, row, 1, f'DAAL floor {floor:.2f}', font=BOLD)
        row += 1
        for j, h in enumerate(['Perturbation step'] + order + ['Total']):
            write_cell(ws, row, 1 + j, h, font=BOLD)
        for _, r in rob[np.isclose(rob['floor'], floor)].iterrows():
            row += 1
            write_cell(ws, row, 1, step_label(r['step'], order))
            for j, k in enumerate(order):
                write_cell(ws, row, 2 + j, int(r[k]), fmt='0')
            write_cell(ws, row, 2 + len(order), f'=SUM(B{row}:{L(1 + len(order))}{row})', fmt='0')
    ws.column_dimensions['A'].width = 24
    ws.freeze_panes = 'B1'
wb.save(ROOT / 'Results.xlsx')


# ---------------------------------------------------------------- Manuscript Tables.xlsx
def table_sheet(wb, name, title, header, rows, fmt, groups=None):
    """A table as printed: title, optional row of column-group titles {column: text}, header, body; NaN prints NA."""
    ws = wb.create_sheet(name)
    write_cell(ws, 1, 1, title, font=TITLE)
    top = 3
    if groups:
        for j, text in groups.items():
            write_cell(ws, top, 1 + j, text, font=BOLD)
        top += 1
    for j, h in enumerate(header):
        c = write_cell(ws, top, 1 + j, h, font=BOLD)
        c.border = Border(top=LINE if not groups else None, bottom=LINE)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if isinstance(v, float) and np.isnan(v):
                v = 'NA'
            numeric = not isinstance(v, str) or v.startswith('=')
            c = write_cell(ws, top + 1 + i, 1 + j, v, fmt=fmt if j and numeric else None)
            c.alignment = Alignment(horizontal='left' if j == 0 else 'right')
    for j in range(len(header)):
        ws.cell(top + len(rows), 1 + j).border = Border(bottom=LINE)
    ws.column_dimensions['A'].width = 16


mt = Workbook()
mt.remove(mt.active)
printed = {'NEO': ('Table 1', 'NEO Embedding Factor Structures Using Transformer Averages for Macro Encoding.', 'sent'),
           'HEXACO': ('Table 4', 'HEXACO Embedding Factor Structures Using Transformer Averages for Atomic Encoding.', 'item')}
for study in ['NEO', 'HEXACO']:
    s = STUDIES[study]
    daal, loadings = data[study][:2]
    name, title, enc = printed[study]
    b = solution_block(loadings[enc], 'avg_trans', s['factors'])
    assert list(b.columns) == s['factors'], f'{name}: solution is not fully aligned'
    facets = s['abbrev'] or s['codes']
    table_sheet(mt, name, f'{name}. {title}', ['Facet'] + s['factors'],
                [[fc] + [r2(x) for x in v] for fc, v in zip(facets, b.values)], '0.00')

supp = {'NEO': ('Table S1', 'Table S2'), 'HEXACO': ('Table S3', 'Table S4')}
for study in ['NEO', 'HEXACO']:
    f = STUDIES[study]['factors']
    al = alignment(data[study][0], f)
    s_structure, s_factor = supp[study]
    n, top, last = len(f), 4, 3 + len(TABLE_ROWS)  # body rows of both tables: 4 to last
    # structure recovery = mean of the factor cells of the factor-recovery table
    rows = [[NAMES[sol]] + [f"=AVERAGE('{s_factor}'!{L(2 + i * n)}{r}:{L(1 + (i + 1) * n)}{r})" for i in range(3)]
            for r, sol in zip(range(top, last + 1), TABLE_ROWS)]
    rows.append(['Overall M'] + [f'=AVERAGE({L(2 + i)}{top}:{L(2 + i)}{last})' for i in range(3)])
    table_sheet(mt, s_structure, f'{s_structure}. {study} structure recovery: proportion of factors aligned.',
                ['Model', 'Atomic', 'AtomicR', 'Macro'], rows, '0.00')
    rows = [[NAMES[sol]] + [int(v) for v in al.loc[sol]] + [f'=AVERAGE(B{r}:{L(1 + 3 * n)}{r})']
            for r, sol in zip(range(top, last + 1), TABLE_ROWS)]
    rows.append(['Overall M'] + [f'=AVERAGE({L(2 + j)}{top}:{L(2 + j)}{last})' for j in range(3 * n + 1)])
    header = ['Model'] + [f'{x} ({e})' for e in ['Atomic', 'AtomicR', 'Macro'] for x in f] + ['M']
    table_sheet(mt, s_factor, f'{s_factor}. {study} factor recovery (1 = aligned).', header, rows, '0.00')
titles = {'NEO': ('Table 2', 'Congruence Between NEO Empirical and NEO Embedding Transformer Models.', '0.00',
                   'Table 3', 'Empirical Correlations Between NEO Empirical and NEO Embedding Transformer Models.'),
          'HEXACO': ('Table 5', 'Congruence Between HEXACO Empirical and HEXACO Embedding Transformer Models.', '.00',
                     'Table 6', 'Correlations Between HEXACO Empirical and HEXACO Embedding Transformer Models.')}
for study in ['NEO', 'HEXACO']:
    f = STUDIES[study]['factors']
    cong, corr = data[study][2], data[study][3]
    t_cong, title_cong, fmt, t_corr, title_corr = titles[study]
    n = len(f)
    table_sheet(mt, t_cong, f'{t_cong}. {title_cong}', ['Model'] + f * 3 + ['M'],
                [[name] + congruence_row(name, v, 5 + i, 5, 2) for i, (name, v) in enumerate(cong.iterrows())], fmt,
                groups={1: 'Atomic congruence', 1 + n: 'Atomic reverse congruence', 1 + 2 * n: 'Macro congruence',
                        1 + 3 * n: 'M'})
    table_sheet(mt, t_corr, f'{t_corr}. {title_corr}', ['Empirical'] + printed_labels(corr.columns, f),
                [[name] + [r2(x) for x in v] for name, v in corr.iterrows()], '.00',
                groups={1: 'Atomic correlations', 1 + n: 'Atomic reverse correlations', 1 + 2 * n: 'Macro correlations'})
for study, name, fmt in [('NEO', 'Table 7', '0.00'), ('HEXACO', 'Table 8', '0')]:
    rob = data[study][4]
    order = list(STUDIES[study]['perturb'])
    rob = rob[np.isclose(rob['floor'], 0.35)]  # the floor used in the paper's tables
    table_sheet(mt, name, f'{name}. {study} Robustness Analysis.', ['Perturbation step'] + order + ['R'],
                [[step_label(r['step'], order)] + [int(r[k]) for k in order] + [f'=SUM(B{4 + i}:{L(1 + len(order))}{4 + i})']
                 for i, (_, r) in enumerate(rob.iterrows())],
                fmt)
# Numbers quoted in the text. Alignment rates are formulas over the supplementary tables; the factor names and
# encodings they refer to, the residuals and the robustness counts come from the results files.
transformers = ['distilroberta', 'miniLM', 'mpnet', 't5', 'psych']  # individual transformers: rows 4-8
text = []
for study in ['NEO', 'HEXACO']:
    f = STUDIES[study]['factors']
    n = len(f)
    s_structure, s_factor = supp[study]
    al = alignment(data[study][0], f)
    S, F = f"'{s_structure}'!", f"'{s_factor}'!"
    pct = lambda ref: f'TEXT({ref},"0%")'
    om = 4 + len(TABLE_ROWS)  # Overall M row
    text.append((study, 'Mean proportion of factors aligned (Overall M of the structure-recovery table), atomic / '
                 'atomic reversed / macro', '=' + '&" / "&'.join(pct(f'{S}{c}{om}') for c in 'BCD'), ''))
    text.append((study, 'Best individual transformer, alignment rate per encoding (atomic / atomic reversed / macro)',
                 '=' + '&" / "&'.join(pct(f'MAX({S}{c}4:{c}8)') for c in 'BCD'), ''))
    text.append((study, 'Worst individual transformer, lowest alignment rate in any encoding',
                 '=' + pct(f'MIN({S}B4:D8)'), ''))
    text.append((study, 'USE-DAN alignment rates, range over encodings',
                 '=' + pct(f'MIN({S}B9:D9)') + '&" to "&' + pct(f'MAX({S}B9:D9)'), ''))
    rng = f'{F}B{om}:{L(1 + 3 * n)}{om}'  # specific-factor rates over the 7 solutions
    factor_rates = al.mean()
    where = [f'{k} ({ENCODINGS[e].lower()})' for (e, k), v in factor_rates.items() if np.isclose(v, factor_rates.min())]
    text.append((study, 'Lowest alignment rate for a specific factor (over the 7 solutions)',
                 '=' + pct(f'MIN({rng})'), ', '.join(where)))
    text.append((study, 'All other specific-factor alignment rates at least',
                 '=' + pct(f'SMALL({rng},COUNTIF({rng},MIN({rng}))+1)'), ''))
    by_enc = {e: al[e].values.mean() for e in ENCODINGS}
    text.append((study, 'Encoding(s) with the highest mean alignment rate', '',
                 ', '.join(ENCODINGS[e] for e in ENCODINGS if np.isclose(by_enc[e], max(by_enc.values())))))
    res = pd.read_csv(ROOT / study / 'results' / 'residuals.csv').iloc[0]
    text.append((study, 'Residuals of the printed solution: RMSR / median |residual| / % over .05 / % over .10',
                 f"{res['rmsr']:.3f} / {res['median_abs']:.3f} / {res['pct_abs_over_05']:.1f}% / "
                 f"{res['pct_abs_over_10']:.1f}%", 'residuals.csv'))
    rob = data[study][4]
    last = rob['step'].iloc[-1]
    text.append((study, 'Robustness: factors aligned in the all-noise matrix at DAAL floor .30 / .35',
                 ' / '.join(str(int(rob[(rob['step'] == last) & np.isclose(rob['floor'], fl)]['total'].iloc[0]))
                            for fl in (0.30, 0.35)), 'robustness.csv'))
ws = mt.create_sheet('Text values')
write_cell(ws, 1, 1, 'Numbers quoted in the text', font=TITLE)
for j, h in enumerate(['Study', 'Quantity', 'Value', 'Which / source']):
    write_cell(ws, 3, 1 + j, h, font=BOLD).border = Border(bottom=LINE)
for i, r in enumerate(text):
    for j, v in enumerate(r):
        write_cell(ws, 4 + i, 1 + j, v)
ws.column_dimensions['B'].width = 95
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 30
order = ['Table 1', 'Table 2', 'Table 3', 'Table 4', 'Table 5', 'Table 6', 'Table 7', 'Table 8',
         'Table S1', 'Table S2', 'Table S3', 'Table S4', 'Text values']
mt._sheets.sort(key=lambda ws: order.index(ws.title))
mt.save(ROOT / 'Manuscript Tables.xlsx')
print('wrote Results.xlsx and Manuscript Tables.xlsx')
